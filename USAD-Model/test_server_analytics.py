"""Regression tests for analytics paths used by packaged builds."""

import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import config
import server
from event_logger import EventLogger


class ServerAnalyticsTests(unittest.TestCase):
    def test_read_csv_uses_configured_runtime_log_directory(self):
        with tempfile.TemporaryDirectory() as runtime_dir, tempfile.TemporaryDirectory() as bundle_dir:
            log_path = Path(runtime_dir) / "violations.csv"
            with log_path.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=["lane", "violation_type"])
                writer.writeheader()
                writer.writerow({"lane": "LANE3", "violation_type": "RED_LIGHT_VIOLATION"})

            # Simulate PyInstaller: server source lives in a temporary bundle,
            # while writable analytics data lives beside the executable.
            with (
                mock.patch.object(config, "LOG_DIRECTORY", runtime_dir),
                mock.patch.object(server, "_THIS_DIR", Path(bundle_dir)),
            ):
                rows = server._read_csv("violations.csv")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["lane"], "LANE3")

    def test_read_csv_returns_empty_list_for_missing_log(self):
        with tempfile.TemporaryDirectory() as runtime_dir:
            with mock.patch.object(config, "LOG_DIRECTORY", runtime_dir):
                self.assertEqual(server._read_csv("accidents.csv"), [])

    def test_log_export_builds_filtered_xlsx(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as runtime_dir:
            log_path = Path(runtime_dir) / "violations.csv"
            with log_path.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=["lane", "license_plate"])
                writer.writeheader()
                writer.writerow({"lane": "LANE3", "license_plate": "ABC123"})
                writer.writerow({"lane": "LANE1", "license_plate": "=HYPERLINK(1)"})

            with mock.patch.object(config, "LOG_DIRECTORY", runtime_dir):
                full = server._build_log_workbook("violations.csv", "Violation Logs")
                filtered = server._build_log_workbook("violations.csv", "Violation Logs", "lane3")

        ws = load_workbook(io.BytesIO(full)).active
        self.assertEqual([c.value for c in ws[1]], ["Lane", "License Plate"])
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws["B3"].data_type, "s")

        ws = load_workbook(io.BytesIO(filtered)).active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["A2"].value, "LANE3")

    def test_log_export_of_missing_log_is_an_empty_workbook(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as runtime_dir:
            with mock.patch.object(config, "LOG_DIRECTORY", runtime_dir):
                data = server._build_log_workbook("accidents.csv", "Accident Logs")
        self.assertEqual(load_workbook(io.BytesIO(data)).active.title, "Accident Logs")

    def test_analytics_report_has_summary_charts_and_log_sheets(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as runtime_dir:
            with (Path(runtime_dir) / "violations.csv").open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=["timestamp", "lane", "violation_type"])
                writer.writeheader()
                writer.writerow({"timestamp": "2026-10-23T14:05:00", "lane": "LANE3",
                                 "violation_type": "RED_LIGHT_VIOLATION"})
                writer.writerow({"timestamp": "2026-10-23T14:30:00", "lane": "LANE3",
                                 "violation_type": "RED_LIGHT_VIOLATION"})
            with (Path(runtime_dir) / "accidents.csv").open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=["timestamp", "duration", "emergency_notified"])
                writer.writeheader()
                writer.writerow({"timestamp": "2026-10-23T09:00:00", "duration": "2.5",
                                 "emergency_notified": "Yes"})

            with mock.patch.object(config, "LOG_DIRECTORY", runtime_dir):
                data = server._build_analytics_workbook()

        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(
            wb.sheetnames,
            ["Summary", "Violation Logs", "Accident Logs", "Traffic Event Logs", "License Plate Logs"],
        )
        summary = wb["Summary"]
        values = {
            row[0]: row[1]
            for row in summary.iter_rows(min_col=1, max_col=3, values_only=True)
            if row[0]
        }
        self.assertEqual(values["Total Violations"], 2)
        self.assertEqual(values["Total Accidents"], 1)
        self.assertEqual(values["Emergency Alerts"], 1)
        self.assertEqual(values["Traffic Events"], 0)
        self.assertEqual(values["14:00"], 2)
        self.assertEqual(values[config.LANES["LANE3"]["name"]], 2)
        self.assertEqual(values["RED LIGHT VIOLATION"], 2)
        self.assertEqual(values["1-5s"], 1)
        self.assertEqual(wb["Violation Logs"].max_row, 3)

    def test_packaged_app_saves_traffic_light_zones(self):
        import asyncio
        import json
        import sys

        original = list(config.TRAFFIC_LIGHT_ZONES)
        try:
            with tempfile.TemporaryDirectory() as exe_dir:
                with (
                    mock.patch.object(sys, "frozen", True, create=True),
                    mock.patch.object(sys, "executable", str(Path(exe_dir) / "USAD.exe")),
                ):
                    result = asyncio.run(server.save_lane_config(
                        {"traffic_light_zones": [[100, 50.4, 46, 15], [1270, 700, 50, 50]]}
                    ))
                saved = json.loads((Path(exe_dir) / "lane_config.json").read_text(encoding="utf-8"))

            self.assertTrue(result["ok"])
            # Out-of-frame zones are clamped to the 1280x720 camera frame.
            self.assertEqual(saved["traffic_light_zones"], [[100, 50, 46, 15], [1270, 700, 10, 20]])
            self.assertEqual(config.TRAFFIC_LIGHT_ZONES, [(100, 50, 46, 15), (1270, 700, 10, 20)])
        finally:
            config.TRAFFIC_LIGHT_ZONES = original

    def _car_size_snapshot(self):
        return (config.CAR_LENGTH_RANGE_PX, config.CAR_WIDTH_RANGE_PX, config.CAR_SIZE_TOLERANCE)

    def _restore_car_size(self, snapshot):
        config.CAR_LENGTH_RANGE_PX, config.CAR_WIDTH_RANGE_PX, config.CAR_SIZE_TOLERANCE = snapshot

    def test_packaged_app_saves_and_reloads_car_size(self):
        import asyncio
        import json
        import sys

        snapshot = self._car_size_snapshot()
        try:
            with tempfile.TemporaryDirectory() as exe_dir:
                json_path = str(Path(exe_dir) / "detection_config.json")
                with (
                    mock.patch.object(sys, "frozen", True, create=True),
                    mock.patch.object(config, "DETECTION_CONFIG_PATH", json_path),
                ):
                    result = asyncio.run(server.save_car_size(
                        {"length_range": [90, 160], "width_range": [50, 95], "tolerance": 0.2}
                    ))
                    self.assertTrue(result["ok"])
                    self.assertEqual(config.CAR_LENGTH_RANGE_PX, (90, 160))
                    saved = json.loads(Path(json_path).read_text(encoding="utf-8"))
                    self.assertEqual(saved["car_width_range_px"], [50, 95])

                    # Simulate the next EXE launch reading the saved file.
                    self._restore_car_size(snapshot)
                    config._load_detection_overrides()
                    self.assertEqual(config.CAR_LENGTH_RANGE_PX, (90, 160))
                    self.assertEqual(config.CAR_WIDTH_RANGE_PX, (50, 95))
                    self.assertEqual(config.CAR_SIZE_TOLERANCE, 0.2)
        finally:
            self._restore_car_size(snapshot)

    def test_source_run_saves_car_size_into_config_py(self):
        import asyncio
        import shutil

        snapshot = self._car_size_snapshot()
        try:
            with tempfile.TemporaryDirectory() as src_dir:
                shutil.copy(Path(config.__file__), Path(src_dir) / "config.py")
                with mock.patch.object(server, "_THIS_DIR", Path(src_dir)):
                    result = asyncio.run(server.save_car_size(
                        {"length_range": [70, 150], "width_range": [40, 85], "tolerance": 0.3}
                    ))
                text = (Path(src_dir) / "config.py").read_text(encoding="utf-8")
            self.assertTrue(result["ok"])
            self.assertIn("\nCAR_LENGTH_RANGE_PX = (70, 150)\n", text)
            self.assertIn("\nCAR_WIDTH_RANGE_PX = (40, 85)\n", text)
            self.assertIn("\nCAR_SIZE_TOLERANCE = 0.3\n", text)
            # The loader inside _load_detection_overrides must stay untouched.
            self.assertIn("CAR_LENGTH_RANGE_PX = tuple(int(v) for v in saved", text)
        finally:
            self._restore_car_size(snapshot)

    def test_invalid_car_size_is_rejected(self):
        import asyncio

        for payload in (
            {"length_range": [150, 90], "width_range": [50, 95]},  # min > max
            {"length_range": [90, 100], "width_range": [50, 120]},  # wider than long
            {"length_range": [90, "x"], "width_range": [50, 95]},
        ):
            result = asyncio.run(server.save_car_size(payload))
            self.assertEqual(result.status_code, 400)

    def test_invalid_traffic_light_zone_is_rejected(self):
        import asyncio

        result = asyncio.run(server.save_lane_config({"traffic_light_zones": [["a", 1, 2, 3]]}))
        self.assertEqual(result.status_code, 400)

    def test_event_logger_preserves_rows_across_restart(self):
        with tempfile.TemporaryDirectory() as runtime_dir:
            with mock.patch.object(config, "LOG_DIRECTORY", runtime_dir):
                logger = EventLogger()
                with open(logger.violation_log_path, "a", newline="", encoding="utf-8") as stream:
                    csv.writer(stream).writerow([
                        "2026-09-13T14:00:00", "2026-09-13", "14:00:00", "14",
                        "RED_LIGHT_VIOLATION", "1", "7", "MEDIUM", "N/A",
                        "LANE3", "RED", "0", "100", "200",
                    ])

                EventLogger()  # Simulate closing and reopening the packaged app.
                with open(logger.violation_log_path, newline="", encoding="utf-8") as stream:
                    rows = list(csv.DictReader(stream))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["lane"], "LANE3")


if __name__ == "__main__":
    unittest.main()
